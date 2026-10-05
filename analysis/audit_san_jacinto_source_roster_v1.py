from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import urllib.request
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

CAPTURE_URL="https://ndownloader.figshare.com/files/33058799"
CAPTURE_SHA="ec70b40fdcc66a3f9c07a3fda64b5eda06d251a899b9bc99cc5c5dff8c4ab301"
META_URL="https://ndownloader.figshare.com/files/33060218"

FOCAL=("CHFA","DKR","LAPM","PEER","PEMA","SKR")

def fetch(url:str)->bytes:
    req=urllib.request.Request(url,headers={"User-Agent":"neon-san-jacinto-roster-audit/1.0"})
    with urllib.request.urlopen(req,timeout=180) as r:
        return r.read()

def clean(x): return str(x or "").strip()

def parse_date(x):
    s=clean(x)
    for fmt in ("%m/%d/%Y","%m/%d/%y","%Y-%m-%d","%m-%d-%Y","%m-%d-%y"):
        try: return datetime.strptime(s,fmt)
        except ValueError: pass
    return None

def season(m):
    if m in (8,9,10): return "fall"
    if m in (11,12,1): return "winter"
    if m in (2,3,4): return "spring"
    if m in (5,6,7): return "summer"
    raise ValueError(m)

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--output",type=Path,required=True)
    args=ap.parse_args()

    raw=fetch(CAPTURE_URL)
    if hashlib.sha256(raw).hexdigest()!=CAPTURE_SHA:
        raise RuntimeError("capture checksum mismatch")
    rows=list(csv.DictReader(raw.decode("utf-8-sig").splitlines()))

    roster=defaultdict(lambda:defaultdict(int))
    ids=defaultdict(set)
    history=defaultdict(lambda:defaultdict(int))
    flags=defaultdict(set)
    for r in rows:
        sp=clean(r.get("species")).upper()
        grid=clean(r.get("grid"))
        dt=parse_date(r.get("date"))
        if not sp or not grid or dt is None: continue
        ss=season(dt.month)
        key=f"{grid}|{ss}"
        roster[key][sp]+=1
        uid=clean(r.get("unique_ID"))
        if uid: ids[sp].add(uid)
        history[sp][clean(r.get("history"))]+=1
        flag=clean(r.get("flag")).upper()
        if flag: flags[key].add(flag)

    all_units=[]
    for g in map(str,range(1,9)):
        for ss in ("fall","winter","spring","summer"):
            key=f"{g}|{ss}"
            counts=dict(sorted(roster[key].items()))
            present=sorted(counts)
            focal_present=[sp for sp in FOCAL if counts.get(sp,0)>0]
            all_units.append({
                "id":key,
                "all_species_present":present,
                "all_species_count":len(present),
                "focal_species_present":focal_present,
                "focal_species_count":len(focal_present),
                "capture_rows_by_species":counts,
                "distinct_nonempty_flags":len(flags[key]),
            })

    meta_raw=fetch(META_URL)
    wb=load_workbook(io.BytesIO(meta_raw),data_only=True,read_only=True)
    sheets={}
    for ws in wb.worksheets:
        vals=[]
        for row in ws.iter_rows(values_only=True):
            rr=[x for x in row if x is not None]
            if rr:
                vals.append([str(x) for x in rr])
        sheets[ws.title]=vals[:200]

    out={
        "schema":"neon.san_jacinto_source_roster_audit.v1",
        "status":"source_semantics_only_no_new_ecological_endpoint",
        "capture_rows":len(rows),
        "species_row_counts":dict(sorted((sp,sum(v.values())) for sp,v in history.items())),
        "unique_id_values_by_species":dict(sorted((sp,len(v)) for sp,v in ids.items())),
        "history_value_counts_by_species":{
            sp:dict(sorted(v.items())) for sp,v in sorted(history.items())
        },
        "grid_season_roster":all_units,
        "metadata_workbook":sheets,
    }
    args.output.parent.mkdir(parents=True,exist_ok=True)
    args.output.write_text(json.dumps(out,indent=2,sort_keys=True)+"\n",encoding="utf-8")
    print(json.dumps(out,indent=2,sort_keys=True))

if __name__=="__main__":
    main()
