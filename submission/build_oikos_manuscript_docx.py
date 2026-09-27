from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "manuscript" / "neon_metacommunity_redundancy" / "MANUSCRIPT_V6_MECHANISM_VALIDATED_DRAFT.md"
DIST = ROOT / "dist"
OUT = DIST / "oikos_main_text_anonymous.docx"


def set_double_spacing(paragraph) -> None:
    paragraph.paragraph_format.line_spacing = 2
    paragraph.paragraph_format.space_after = Pt(0)


def add_page_number(paragraph) -> None:
    paragraph.alignment = 2
    run = paragraph.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def add_continuous_line_numbers(section) -> None:
    sect_pr = section._sectPr
    old = sect_pr.find(qn("w:lnNumType"))
    if old is not None:
        sect_pr.remove(old)
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")
    ln.set(qn("w:distance"), "360")
    sect_pr.append(ln)


def add_inline_runs(paragraph, text: str) -> None:
    # Minimal Markdown support for bold and italics used by the manuscript.
    pattern = re.compile(r"(\*\*.*?\*\*|\*.*?\*)")
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        token = m.group(0)
        if token.startswith("**"):
            r = paragraph.add_run(token[2:-2])
            r.bold = True
        else:
            r = paragraph.add_run(token[1:-1])
            r.italic = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def build() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()

    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.title = "Anonymous main text"
    props.subject = ""
    props.keywords = ""
    props.comments = ""

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)

    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"

    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    add_continuous_line_numbers(section)
    add_page_number(section.footer.paragraphs[0])

    in_abstract = False
    page_break_inserted = False

    for raw in lines:
        line = raw.rstrip()

        if not line:
            p = doc.add_paragraph()
            set_double_spacing(p)
            continue

        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            add_inline_runs(p, line[2:].strip())
            set_double_spacing(p)
            continue

        if line == "## Abstract":
            p = doc.add_paragraph(style="Heading 1")
            p.add_run("Abstract")
            set_double_spacing(p)
            in_abstract = True
            continue

        if line.startswith("## 1. Introduction"):
            if not page_break_inserted:
                pbreak = doc.add_paragraph()
                pbreak.add_run().add_break(WD_BREAK.PAGE)
                page_break_inserted = True
            in_abstract = False

        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline_runs(p, line[4:].strip())
            set_double_spacing(p)
            continue

        if line.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            # Strip numeric section prefix from style text only if present? Keep as manuscript text.
            add_inline_runs(p, line[3:].strip())
            set_double_spacing(p)
            continue

        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, line[2:].strip())
            set_double_spacing(p)
            continue

        if re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph(style="List Number")
            add_inline_runs(p, re.sub(r"^\d+\.\s+", "", line))
            set_double_spacing(p)
            continue

        p = doc.add_paragraph()
        add_inline_runs(p, line)
        set_double_spacing(p)

    DIST.mkdir(exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
